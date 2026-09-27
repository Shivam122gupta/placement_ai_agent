import io
import pytest
from httpx import AsyncClient
from docx import Document
from app.models.user import UserDocument


def create_test_docx():
    doc = Document()
    doc.add_heading("Ankit Sharma", level=1)
    doc.add_paragraph("Email: ankit.developer@example.com | Phone: +91 9876543210 | Bengaluru, India")
    doc.add_paragraph("Summary: Proactive Backend Engineer with 1 year experience building scalable FastAPI microservices.")
    doc.add_paragraph("Education: B.Tech Computer Science and Engineering from National Institute of Technology, 2025, CGPA: 8.8")
    doc.add_paragraph("Experience: Backend Intern at Stripe. Built distributed payment webhook pipeline and optimized database.")
    doc.add_paragraph("Skills: Python, FastAPI, MongoDB, PostgreSQL, Docker, Redis, Kubernetes")
    doc.add_paragraph("Project: AI Placement Agent - Autonomous career platform with vector search and multi-agent reasoning.")
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


@pytest.mark.asyncio
async def test_auto_profile_population_on_resume_upload(
    async_client: AsyncClient, auth_headers: dict, test_user: UserDocument
):
    # Upload valid docx resume
    file_stream = create_test_docx()
    files = {
        "file": ("ankit_sharma_resume.docx", file_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    }

    upload_res = await async_client.post("/api/v1/resumes", files=files, headers=auth_headers)
    assert upload_res.status_code == 201
    resume_data = upload_res.json()["data"]
    assert resume_data["status"] == "COMPLETED"

    # Verify Profile was auto-populated immediately without manual typing!
    profile_res = await async_client.get("/api/v1/profile", headers=auth_headers)
    assert profile_res.status_code == 200
    pdata = profile_res.json()["data"]

    assert pdata["full_name"] == "Ankit Sharma"
    assert len(pdata["skills"]) >= 3
    assert len(pdata["education"]) >= 1
    assert pdata["completion_score"] >= 70


@pytest.mark.asyncio
async def test_experience_sub_resource_crud(async_client: AsyncClient, auth_headers: dict):
    # 1. Add Experience
    exp_payload = {
        "company": "Stripe",
        "role": "Backend Intern",
        "duration": "June 2025 - August 2025",
        "location": "Remote",
        "highlights": ["Built distributed webhook pipeline with FastAPI", "Optimized MongoDB aggregations"],
    }
    create_res = await async_client.post("/api/v1/profile/experience", json=exp_payload, headers=auth_headers)
    assert create_res.status_code == 201
    profile = create_res.json()["data"]
    assert len(profile["experience"]) >= 1
    exp_id = profile["experience"][0]["id"]

    # 2. Delete Experience
    del_res = await async_client.delete(f"/api/v1/profile/experience/{exp_id}", headers=auth_headers)
    assert del_res.status_code == 200
    updated_profile = del_res.json()["data"]
    assert not any(e["id"] == exp_id for e in updated_profile["experience"])


@pytest.mark.asyncio
async def test_auto_sync_latest_resume_endpoint(async_client: AsyncClient, auth_headers: dict):
    # 1. Upload resume first
    file_stream = create_test_docx()
    files = {
        "file": ("ankit_sharma_resume.docx", file_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    }
    upload_res = await async_client.post("/api/v1/resumes", files=files, headers=auth_headers)
    assert upload_res.status_code == 201

    # 2. Call 1-click auto sync
    sync_res = await async_client.post("/api/v1/profile/auto-sync-latest-resume", headers=auth_headers)
    assert sync_res.status_code == 200
    pdata = sync_res.json()["data"]
    assert pdata["full_name"] == "Ankit Sharma"
    assert pdata["completion_score"] >= 70
