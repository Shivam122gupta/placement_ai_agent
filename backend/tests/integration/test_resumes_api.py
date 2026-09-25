import io
import pytest
from httpx import AsyncClient
from docx import Document


@pytest.mark.asyncio
async def test_resume_upload_parse_sync_and_delete_flow(async_client: AsyncClient, auth_headers: dict):
    # 1. Create a dummy test docx resume in-memory
    doc = Document()
    doc.add_heading("Ankit Sharma", level=1)
    doc.add_paragraph("Email: ankit@example.com | Bengaluru, India")
    doc.add_paragraph("Education: B.Tech Computer Science from National Institute of Technology, 2025")
    doc.add_paragraph("Skills: Python, FastAPI, MongoDB, Docker, Machine Learning")
    doc.add_paragraph("Project: AI Placement Agent - Multi-agent career platform")
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    # 2. Upload resume via POST /api/v1/resumes
    files = {
        "file": ("ankit_resume.docx", file_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    }
    upload_res = await async_client.post("/api/v1/resumes", files=files, headers=auth_headers)
    assert upload_res.status_code == 201
    upload_json = upload_res.json()
    assert upload_json["success"] is True
    assert upload_json["data"]["filename"] == "ankit_resume.docx"
    assert upload_json["data"]["status"] == "COMPLETED"
    assert "skills" in upload_json["data"]["parsed_data"]
    resume_id = upload_json["data"]["id"]

    # 3. List candidate resumes
    list_res = await async_client.get("/api/v1/resumes", headers=auth_headers)
    assert list_res.status_code == 200
    assert len(list_res.json()["data"]) >= 1

    # 4. Get single resume details
    get_res = await async_client.get(f"/api/v1/resumes/{resume_id}", headers=auth_headers)
    assert get_res.status_code == 200
    assert get_res.json()["data"]["id"] == resume_id

    # 5. Download original resume
    download_res = await async_client.get(f"/api/v1/resumes/{resume_id}/download", headers=auth_headers)
    assert download_res.status_code == 200
    assert len(download_res.content) > 0

    # 6. Sync parsed data to active Candidate Profile
    sync_payload = {
        "sync_personal": True,
        "sync_education": True,
        "sync_skills": True,
        "sync_projects": True,
        "sync_certifications": True,
    }
    sync_res = await async_client.post(f"/api/v1/resumes/{resume_id}/sync-profile", json=sync_payload, headers=auth_headers)
    assert sync_res.status_code == 200
    synced_profile = sync_res.json()["data"]
    assert len(synced_profile["skills"]) >= 3
    assert synced_profile["completion_score"] > 50

    # 7. Delete resume
    del_res = await async_client.delete(f"/api/v1/resumes/{resume_id}", headers=auth_headers)
    assert del_res.status_code == 200

    # 8. Confirm deletion
    get_after_del = await async_client.get(f"/api/v1/resumes/{resume_id}", headers=auth_headers)
    assert get_after_del.status_code == 404
