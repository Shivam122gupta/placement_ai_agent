import logging
from pathlib import Path
from docx import Document
from app.services.parsers.base import BaseDocumentParser
from app.core.exceptions import ValidationError

logger = logging.getLogger("app.parsers.docx")


class DOCXParser(BaseDocumentParser):
    def extract_text(self, file_path: Path) -> str:
        try:
            doc = Document(str(file_path))
            full_text = []

            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)

            # Extract tables content
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))

            cleaned = self.clean_text("\n\n".join(full_text))
            return cleaned
        except Exception as e:
            logger.error(f"Failed to extract text from DOCX {file_path}: {e}")
            raise ValidationError(f"Could not parse Word DOCX file: {str(e)}", code="DOCX_PARSING_FAILED")
